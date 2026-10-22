import json
from pathlib import Path
import re
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/handoff_20261021'
MD = OUT / '社区老人健康风险评估系统部署与数据交接.md'
source = MD.read_text(encoding='utf-8')
manifest = json.loads((OUT / 'database/database_manifest.json').read_text(encoding='utf-8'))
rows = ['| 数据表 | 记录数 | 字段数 |', '| --- | --- | --- |']
rows += [f"| {t['table']} | {t['rows']} | {t['columns']} |" for t in manifest['tables']]
source = source.replace('{{DATABASE_TABLE}}', '\n'.join(rows))
MD.write_text(source, encoding='utf-8')

doc = Document()
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin, sec.bottom_margin = Mm(17), Mm(25)
sec.left_margin, sec.right_margin = Mm(19), Mm(19)
sec.footer_distance = Mm(8)


def font(style, size, bold=False):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '微软雅黑')


for name, size, bold in [('Normal',10.5,False), ('Title',22,True), ('Subtitle',11,False), ('Heading 1',16,True), ('Heading 2',11.5,True)]:
    font(doc.styles[name], size, bold)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.16
normal.space_after = Pt(5)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_before = Pt(0)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(5)
for h in ['Heading 1','Heading 2']:
    doc.styles[h].paragraph_format.keep_with_next = True

code = doc.styles.add_style('Command', 1)
font(code, 8.4)
code.font.name = 'Consolas'
code.paragraph_format.left_indent = Mm(2)
code.paragraph_format.line_spacing = 1.05
code.paragraph_format.space_after = Pt(1)
code.paragraph_format.keep_with_next = False

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run('项目部署与数据交接  |  ')
run.font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
footer._p.append(fld)


def table(lines):
    cells = [[v.strip() for v in line.strip().strip('|').split('|')] for line in lines if not re.match(r'^\|\s*:?-+', line)]
    t = doc.add_table(rows=1, cols=len(cells[0]))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    n = len(cells[0])
    if cells[0][0] == '数据表':
        widths = [116,28,28]
    elif n == 5:
        widths = [38,30,32,32,40]
    elif n == 3:
        widths = [86,43,43]
    else:
        widths = [51,121]
    for c,w in zip(t.columns,widths): c.width = Mm(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for name in ['top','left','bottom','right','insideH','insideV']:
        edge=OxmlElement('w:'+name)
        edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9')
        borders.append(edge)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for name,val in [('top',60),('bottom',60),('left',90),('right',90)]:
        e=OxmlElement('w:'+name);e.set(qn('w:w'),str(val));e.set(qn('w:type'),'dxa');margins.append(e)
    props.append(margins)
    for idx,record in enumerate(cells):
        row=t.rows[0] if idx==0 else t.add_row()
        pr=row._tr.get_or_add_trPr()
        pr.append(OxmlElement('w:cantSplit'))
        if idx==0:pr.append(OxmlElement('w:tblHeader'))
        for j,(cell,value) in enumerate(zip(row.cells,record)):
            cell.width=Mm(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.08
            p.alignment=WD_ALIGN_PARAGRAPH.CENTER if (n>=3 and j>0) else WD_ALIGN_PARAGRAPH.LEFT
            r=p.add_run(value);r.font.size=Pt(9 if n>=3 else 9.4);r.bold=idx==0
            if idx==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1');cell._tc.get_or_add_tcPr().append(sh)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(2);p.add_run().font.size=Pt(2)


lines=source.splitlines();i=0;in_code=False
while i<len(lines):
    line=lines[i]
    if line.startswith('```'):
        in_code=not in_code
        if not in_code:
            p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=Pt(2);p.add_run().font.size=Pt(2)
        i+=1;continue
    if in_code:
        p=doc.add_paragraph(line,style='Command')
        i+=1;continue
    if line=='<!-- page -->':
        i+=1;continue
    if line.startswith('|'):
        group=[]
        while i<len(lines) and lines[i].startswith('|'):
            group.append(lines[i]);i+=1
        table(group);continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],style='Title');p.paragraph_format.space_after=Pt(8)
    elif line.startswith('## '):doc.add_paragraph(line[3:],style='Heading 1')
    elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
    elif line.strip():
        p=doc.add_paragraph(line)
        if line.startswith('交接日期'):
            p.style=doc.styles['Subtitle']
    i+=1

doc.core_properties.title='社区老人健康风险评估系统部署与数据交接'
doc.core_properties.subject='Windows部署 MySQL迁移 联合模型运行与交接验收'
doc.core_properties.author='项目交接'
dest=MD.with_suffix('.docx')
doc.save(dest)
print(dest)
